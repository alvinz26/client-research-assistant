"""Demo mode runs the whole workflow offline and produces a clearly labelled Word file."""
import re

from docx import Document

from research.config import Config
from research.models import Brief
from research.pipeline import run


def _demo_config(tmp_path):
    return Config(api_key=None, model="claude-sonnet-5-5", max_searches=8, demo=True, output_dir=tmp_path)


def _doc_text(path):
    doc = Document(path)
    body = [p.text for p in doc.paragraphs]
    footer = [p.text for p in doc.sections[0].footer.paragraphs]
    return body, " ".join(footer)


def test_demo_run_produces_labelled_cited_docx(tmp_path):
    steps = []
    result = run(Brief("Harbourline Facades", email="j.doe@harbourline.example"), _demo_config(tmp_path), steps.append)

    assert result.path.exists() and result.path.suffix == ".docx"
    assert steps, "progress callback should be called"
    body, footer = _doc_text(result.path)
    text = "\n".join(body)
    assert "DEMO" in body[0]
    assert "Sources and evidence" in body
    assert "j.doe@harbourline.example" in text  # input clues are recorded
    assert "AI-assisted" in footer

    # Every bullet in the fact sections carries at least one [n] that exists in the source list.
    numbers = {int(n) for n in re.findall(r"^\[(\d+)\]", text, re.M)}
    for section in result.report.sections:
        for bullet in section.bullets:
            assert bullet.refs and set(bullet.refs) <= numbers
    assert any(s.quotes for s in result.report.sources)


def test_chinese_report_uses_chinese_labels(tmp_path):
    result = run(Brief("Harbourline Facades", language="zh"), _demo_config(tmp_path))
    body, footer = _doc_text(result.path)
    assert "来源与证据" in body
    assert "DEMO" in body[0]
    assert "AI" in footer
